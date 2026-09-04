"""Excel 导出（openpyxl）。候选表 + 评分表 + 邀请表三 sheet。
不导出手机/邮箱（合规红线：只导出最小可操作字段）。"""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import db
from .liepin import adapter as lp

_HEADER_FILL = PatternFill("solid", fgColor="E8F1FB")
_HEADER_FONT = Font(bold=True)


def _write_sheet(ws, headers: list[str], rows: list[list]):
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    for r in rows:
        ws.append(r)
    for i, _h in enumerate(headers, start=1):
        width = min(40, max(10, (max((len(str(r[i - 1])) for r in rows), default=10) + 2)))
        ws.column_dimensions[get_column_letter(i)].width = width


def build_run_export(run_id: int) -> bytes:
    """生成一个 run 的 xlsx 字节流：候选 / 评分 / 邀请 三 sheet。"""
    wb = Workbook()
    cands = db.list_run_candidates(run_id)
    cands = db.attach_scores(cands)
    run = db.get_run(run_id) or {}

    # Sheet1 候选人
    ws = wb.active
    ws.title = "候选人"
    headers = ["姓名", "期望职位", "期望薪资", "当前公司", "工作年限", "学历", "学校",
               "年龄", "现居地", "活跃度", "状态", "规则评分", "是否新发现", "简历链接"]
    rows = [[c.get("name"), c.get("desired_title"), c.get("desired_salary"),
             c.get("current_company"), c.get("work_years"), c.get("edu"), c.get("school"),
             c.get("age"), c.get("current_city"), c.get("active_status"),
             "已通过" if c.get("hard_gate") else "未通过",
             c.get("score", {}).get("total") if c.get("score") else None,
             "新" if c.get("is_new") else "已有",
             lp.resume_detail_url(c.get("resume_id")) or ""]
            for c in cands]
    _write_sheet(ws, headers, rows)

    # Sheet2 评分明细（每个候选一行，含五维）
    ws2 = wb.create_sheet("评分明细")
    headers2 = ["姓名", "总评分", "专业能力", "技术领导力", "经营思维", "资源整合",
                "岗位契合度", "附加分", "评分方法", "说明"]
    rows2 = []
    for c in cands:
        sc = c.get("score") or {}
        dims = sc.get("dims") or {}
        rows2.append([c.get("name"), sc.get("total"), dims.get("专业能力"),
                      dims.get("技术领导力"), dims.get("经营思维"), dims.get("资源整合"),
                      dims.get("岗位契合度"), sc.get("bonus"),
                      sc.get("method", "rule"), sc.get("source", "")])
    _write_sheet(ws2, headers2, rows2)

    # Sheet3 邀请记录
    ws3 = wb.create_sheet("邀请")
    invites = db.list_invites(run_id)
    headers3 = ["姓名", "岗位 ejobId", "状态", "错误", "发送时间"]
    rows3 = [[i.get("candidate_name"), i.get("job_id"), i.get("status"), i.get("error"),
              i.get("sent_at")] for i in invites]
    _write_sheet(ws3, headers3, rows3)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
