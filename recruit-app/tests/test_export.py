"""export：build_run_export 产出合法 xlsx（三 sheet），且不含敏感字段。"""
from io import BytesIO

from openpyxl import load_workbook

from backend import db
from backend.export import build_run_export


def _seed_run():
    rid = db.create_search_run(0, "manual")
    cid, _ = db.upsert_candidate(rid, {
        "resume_id": "EXP1", "name": "李四", "desired_title": "研发总监",
        "desired_salary": "70-90K", "current_company": "华润微", "current_city": "厦门",
        "work_years": "14年", "edu": "博士", "school": "清华大学", "age": 42,
        "active_status": "在职", "resume_url": "https://lpt.liepin.com/cv/exp1",
    })
    db.update_candidate_gate(cid, 1, "", 1, "命中目标公司")
    db.save_score(cid, "llm", {"专业能力": 20, "技术领导力": 18}, 92.5, "deepseek-chat",
                  {"附加分": "+5"})
    db.create_invite(cid, rid, "J-EJOB-1")
    db.set_invite_status(db.list_invites(rid)[0]["id"], "sent")
    return rid


def test_export_returns_valid_xlsx_bytes():
    rid = _seed_run()
    data = build_run_export(rid)
    assert isinstance(data, bytes)
    assert data[:2] == b"PK"                     # xlsx 是 zip 容器
    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames == ["候选人", "评分明细", "邀请"]


def test_export_sheets_content():
    rid = _seed_run()
    wb = load_workbook(BytesIO(build_run_export(rid)))

    cand = list(wb["候选人"].values)
    assert cand[0][0] == "姓名"
    assert cand[1][0] == "李四"
    assert cand[1][3] == "华润微"                 # 当前公司
    assert cand[1][11] == 92.5                    # 规则评分列取 LLM total

    score = list(wb["评分明细"].values)
    assert score[0][1] == "总评分"
    assert score[1][1] == 92.5
    assert score[1][9] == "deepseek-chat"         # 评分方法 source

    inv = list(wb["邀请"].values)
    assert inv[0][2] == "状态"
    assert inv[1][1] == "J-EJOB-1"
    assert inv[1][2] == "sent"


def test_export_empty_run_still_valid():
    rid = db.create_search_run(0, "manual")
    wb = load_workbook(BytesIO(build_run_export(rid)))
    assert wb.sheetnames == ["候选人", "评分明细", "邀请"]
    assert len(list(wb["候选人"].values)) == 1    # 只有表头


def test_export_never_contains_contact_fields():
    rid = _seed_run()
    blob = build_run_export(rid).decode("utf-8", errors="ignore")
    for banned in ("phone", "mobile", "email", "微信", "手机", "电话", "邮箱"):
        assert banned not in blob
